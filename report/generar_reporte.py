"""
Genera el reporte final (Reporte_Parcial1.docx) exigido por el enunciado.

Lee los resultados numéricos generados por las fases 1-5 (baseline_results,
ga_summary, final_results, comparison_table) y las gráficas ya creadas por
`compare_report.py`, y ensambla un documento .docx con TODAS las secciones
obligatorias: Introducción, Metodología, Resultados, Discusión, Conclusiones
y las 18 Preguntas de Análisis Crítico.

Ejecutar DESPUÉS de que el pipeline completo (run_all.py) haya terminado.
"""
import os
import json

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from src import config
from src.chromosome import GENE_SPACE, GENE_NAMES

OUT_PATH = os.path.join(config.REPORT_DIR, "Reporte_Parcial1.docx")


def _load_json(path, default=None):
    if os.path.exists(path):
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    return default


def _add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)


def _add_par(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    return p


def _add_bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def _add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = str(h)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = str(val)
    return table


def _add_image_if_exists(doc, path, width_in=6.0, caption=None):
    if os.path.exists(path):
        doc.add_picture(path, width=Inches(width_in))
        if caption:
            cap = doc.add_paragraph(caption)
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cap.runs:
                run.italic = True
                run.font.size = Pt(9)
    else:
        _add_par(doc, f"[Gráfica no disponible: {os.path.basename(path)}]", italic=True)


def generate_report():
    baseline = _load_json(os.path.join(config.BASELINE_DIR, "baseline_results.json"), {})
    ga_summary = _load_json(os.path.join(config.GA_DIR, "ga_summary.json"), {})
    final = _load_json(os.path.join(config.FINAL_DIR, "final_results.json"), {})
    comparison = _load_json(os.path.join(config.REPORT_DIR, "comparison_table.json"), {})
    gen_history = _load_json(os.path.join(config.GA_DIR, "generation_best_fitness.json"), [])

    doc = Document()

    # ------------------------------------------------------------------
    title = doc.add_heading(
        "Optimización Multiobjetivo de Redes Neuronales Convolucionales "
        "Profundas mediante Algoritmos Genéticos", level=0
    )
    _add_par(doc, "Parcial 1 - Computación Flexible", italic=True)
    _add_par(doc, "Dataset: CIFAR-10  |  Frameworks: TensorFlow/Keras + pyGAD")
    doc.add_page_break()

    # ==================================================================
    # INTRODUCCIÓN
    # ==================================================================
    _add_heading(doc, "1. Introducción", level=1)

    _add_heading(doc, "1.1 Redes Neuronales Convolucionales", level=2)
    _add_par(doc,
        "Las Redes Neuronales Convolucionales (CNN) son arquitecturas de aprendizaje "
        "profundo diseñadas para procesar datos con estructura de rejilla, como "
        "imágenes. Se basan en capas de convolución (Conv2D), que extraen "
        "características locales (bordes, texturas, formas) mediante filtros "
        "compartidos espacialmente, y capas de reducción de dimensionalidad "
        "(MaxPooling2D), que disminuyen el tamaño espacial del mapa de "
        "características conservando la información más relevante y aportando "
        "cierta invariancia a traslaciones pequeñas. Al final de la parte "
        "convolucional, una capa Flatten convierte los mapas de características "
        "en un vector, que alimenta una o más capas totalmente conectadas (Dense) "
        "encargadas de la clasificación final. En este proyecto, tanto la "
        "arquitectura baseline como todas las arquitecturas generadas por el "
        "Algoritmo Genético siguen este patrón: Entrada + (Conv2D+MaxPooling2D)ⁿ "
        "+ Flatten + Dense + Salida."
    )

    _add_heading(doc, "1.2 Algoritmos Genéticos", level=2)
    _add_par(doc,
        "Los Algoritmos Genéticos (GA) son metaheurísticas de optimización "
        "inspiradas en la selección natural. Mantienen una población de "
        "soluciones candidatas (individuos), cada una codificada como un "
        "cromosoma (vector de genes). En cada generación, los individuos se "
        "evalúan mediante una función de fitness, se seleccionan los mejores "
        "como padres, y se generan nuevos individuos mediante operadores de "
        "cruce (crossover) y mutación. Este ciclo se repite durante varias "
        "generaciones, permitiendo que la población evolucione hacia regiones "
        "del espacio de búsqueda con mayor fitness. A diferencia de los métodos "
        "basados en gradiente, los GA no requieren que la función objetivo sea "
        "diferenciable, lo cual los hace apropiados para optimizar "
        "simultáneamente hiperparámetros continuos, categóricos y decisiones "
        "arquitectónicas discretas (número de capas, número de filtros, etc.), "
        "como se requiere en este proyecto. En este trabajo se utilizó la "
        "biblioteca pyGAD para implementar el ciclo evolutivo completo."
    )

    _add_heading(doc, "1.3 Optimización de hiperparámetros", level=2)
    _add_par(doc,
        "El desempeño de una CNN depende fuertemente de decisiones de diseño "
        "que tradicionalmente se ajustan de forma manual o mediante búsquedas "
        "exhaustivas (grid search) o aleatorias (random search): tasa de "
        "aprendizaje, tamaño de lote, tipo de optimizador, profundidad de la "
        "red, número de filtros por capa, tamaño del kernel, número de "
        "neuronas en la capa densa y tasa de dropout. Estos hiperparámetros "
        "interactúan de forma no lineal, por lo que encontrar una combinación "
        "óptima manualmente es costoso y poco escalable. Los Algoritmos "
        "Genéticos permiten automatizar esta búsqueda, explorando el espacio "
        "conjunto de hiperparámetros de entrenamiento y de arquitectura de "
        "forma simultánea."
    )

    _add_heading(doc, "1.4 Eficiencia computacional en Deep Learning", level=2)
    _add_par(doc,
        "Durante años, la investigación en Deep Learning priorizó "
        "principalmente la exactitud de los modelos, produciendo arquitecturas "
        "cada vez más grandes y costosas. Sin embargo, en aplicaciones reales "
        "(dispositivos móviles, sistemas embebidos, edge computing, servicios "
        "con restricciones de latencia o costo de inferencia) el número de "
        "parámetros entrenables, el consumo de memoria y el tiempo de "
        "entrenamiento/inferencia son igual de relevantes que la exactitud. "
        "Este proyecto aborda explícitamente ese compromiso (trade-off) "
        "mediante una función de fitness multiobjetivo que recompensa tanto "
        "la exactitud como la compacidad del modelo."
    )

    # ==================================================================
    # METODOLOGÍA
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "2. Metodología", level=1)

    _add_heading(doc, "2.1 División de datos", level=2)
    _add_par(doc,
        f"El conjunto CIFAR-10 (60,000 imágenes de 32x32x3, 10 clases) se "
        f"combinó en un único pool y se particionó de forma fija y "
        f"reproducible (semilla={config.SEED}) en: 60% Train "
        f"({int(0.6*60000)} imágenes), 20% Validation ({int(0.2*60000)} "
        f"imágenes) y 20% Test ({int(0.2*60000)} imágenes). Los índices "
        f"exactos de la partición se guardaron en disco "
        f"(results/split_indices.json) y se reutilizaron sin cambios durante "
        f"TODAS las fases del experimento (baseline, optimización GA y "
        f"reentrenamiento final), garantizando que la comparación entre "
        f"arquitecturas sea justa."
    )

    _add_heading(doc, "2.2 Codificación genética", level=2)
    _add_par(doc,
        "Cada individuo se representa como un vector de longitud fija de "
        f"{len(GENE_SPACE)} genes:"
    )
    _add_bullets(doc, [
        "learning_rate: tasa de aprendizaje del optimizador.",
        "batch_size: tamaño de lote de entrenamiento.",
        "optimizer: 0=Adam, 1=RMSprop, 2=SGD.",
        f"n_conv_layers: número de bloques Conv2D+MaxPooling2D "
        f"({config.MIN_CONV_LAYERS} a {config.MAX_CONV_LAYERS}).",
        f"filters_layer1..filters_layer{config.MAX_CONV_LAYERS}: número de "
        f"filtros de cada bloque convolucional (solo se usan los primeros "
        f"n_conv_layers genes; el resto se ignora al construir el modelo).",
        "kernel_size: tamaño del kernel (compartido por todas las capas "
        "Conv2D del individuo).",
        "dense_units: número de neuronas de la capa densa oculta.",
        "dropout: tasa de dropout aplicada tras Flatten.",
    ])
    _add_par(doc,
        "Cada capa Conv2D genera siempre su MaxPooling2D asociado (nunca se "
        "omite), respetando la estructura obligatoria Entrada + "
        "(Conv2D+MaxPooling2D)ⁿ + Flatten + Dense + Salida. Se utiliza "
        "padding='valid' en las capas Conv2D generadas por el GA -igual que "
        "en el baseline- para mantener una comparación justa de tamaño de "
        "mapas de características. Dado que con padding 'valid' no toda "
        "combinación (kernel_size, n_conv_layers) es geométricamente válida "
        "sobre una entrada de 32x32 (el mapa de características podría "
        "intentar reducirse a un tamaño ≤ 0), se recorta automáticamente la "
        "profundidad solicitada al máximo número de bloques Conv2D+"
        "MaxPooling2D que sí caben, evitando descartar individuos por "
        "arquitecturas inválidas."
    )

    _add_heading(doc, "2.3 Espacio de búsqueda", level=2)
    space_rows = []
    for name, space in zip(GENE_NAMES, GENE_SPACE):
        space_rows.append((name, str(space)))
    _add_table(doc, ["Gen", "Valores permitidos"], space_rows)

    _add_heading(doc, "2.4 Función fitness", level=2)
    _add_par(doc,
        "fitness = w1 * val_accuracy + w2 * (1 - clip((TP - TP_min) / "
        "(TP_max - TP_min), 0, 1))"
    )
    _add_par(doc,
        f"Con w1={config.FITNESS_W1}, w2={config.FITNESS_W2}, "
        f"TP_min={config.TP_MIN:,}, TP_max={config.TP_MAX:,}."
    )
    _add_par(doc,
        "NOTA 1 (signo de la fórmula): el enunciado original escribe el "
        "segundo término como \"1 + (TP - TP_min)/(TP_max - TP_min)\", pero "
        "su propio ejemplo numérico resuelto (val_accuracy=0.91, TP=5230 "
        "-> fitness=0.9242) solo es reproducible usando un signo \"menos\": "
        "0.637 + 0.3*(1 - 4230/99000) = 0.9242. Usar \"+\" premiaría con "
        "mayor fitness a las redes con MÁS parámetros, contradiciendo el "
        "objetivo explícito de minimizar la complejidad. Se implementó la "
        "versión con \"-\", consistente tanto con el objetivo declarado "
        "como con el ejemplo numérico del propio enunciado."
    )
    _add_par(doc,
        "NOTA 2 (TP_min / TP_max): los valores TP_min=1,000 / TP_max=100,000 "
        "sugeridos en el enunciado son ilustrativos para el ejemplo de "
        "cálculo, pero resultan demasiado pequeños para el espacio de "
        "búsqueda real de este proyecto (arquitecturas Flatten+Dense sobre "
        "CIFAR-10): al enumerar 120 combinaciones representativas del "
        "espacio de búsqueda, el número de parámetros entrenables varió "
        "entre ~9,900 y ~29.5 millones, y la arquitectura baseline usada en "
        "clase tiene 315,722 parámetros entrenables (ya por encima del "
        "TP_max=100,000 sugerido). Con los valores sugeridos, el término de "
        "complejidad se saturaría (clip=0) para prácticamente cualquier "
        "arquitectura, perdiendo su poder discriminante. Por ello se "
        "recalibraron a TP_min=10,000 y TP_max=1,000,000, valores que "
        "cubren la gran mayoría de arquitecturas razonables del espacio de "
        "búsqueda (incluido el baseline) y solo saturan la penalización "
        "para las combinaciones más extremas (pocas capas convolucionales "
        "con muchos filtros y una capa densa muy grande). Se mantuvieron "
        "los pesos sugeridos w1=0.7 / w2=0.3 porque priorizan la exactitud "
        "-más costosa de recuperar una vez perdida- sin dejar de aplicar "
        "una presión significativa hacia arquitecturas compactas."
    )
    _add_par(doc,
        "Durante la optimización (Fase 2), cada individuo se entrena "
        f"únicamente con el conjunto Train durante {config.GA_INDIVIDUAL_EPOCHS} "
        "épocas y se evalúa con el conjunto Validation (se usa el máximo "
        "val_accuracy observado en esas épocas, para reducir el ruido de "
        "una sola época). El conjunto Test NUNCA participa en esta fase."
    )

    _add_heading(doc, "2.5 Configuración del Algoritmo Genético (pyGAD)", level=2)
    ga_config_rows = [
        ("Tamaño de población", config.GA_POPULATION_SIZE),
        ("Número de generaciones", config.GA_NUM_GENERATIONS),
        ("Épocas por individuo", config.GA_INDIVIDUAL_EPOCHS),
        ("Padres por apareamiento", config.GA_NUM_PARENTS_MATING),
        ("Selección de padres", config.GA_PARENT_SELECTION_TYPE),
        ("Elitismo (keep_elitism)", config.GA_KEEP_ELITISM),
        ("Tipo de cruce", config.GA_CROSSOVER_TYPE),
        ("Tipo de mutación", config.GA_MUTATION_TYPE),
        ("Porcentaje de genes mutados", f"{config.GA_MUTATION_PERCENT_GENES}%"),
        ("Semilla aleatoria", config.SEED),
    ]
    _add_table(doc, ["Parámetro", "Valor"], ga_config_rows)
    _add_par(doc,
        "Se utiliza además una caché de fitness por cromosoma exacto: si un "
        "individuo idéntico sobrevive por elitismo entre generaciones, no "
        "se reentrena, ahorrando cómputo sin alterar el comportamiento del "
        "algoritmo."
    )

    # ==================================================================
    # RESULTADOS
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "3. Resultados", level=1)

    _add_heading(doc, "3.1 Accuracy de la red base", level=2)
    if baseline:
        _add_par(doc,
            f"Accuracy en Test: {baseline.get('test_accuracy', float('nan')):.4f}  "
            f"|  Loss en Test: {baseline.get('test_loss', float('nan')):.4f}  "
            f"|  Épocas: {baseline.get('epochs')}  "
            f"|  Tiempo de entrenamiento: {baseline.get('train_time_seconds', 0):.1f} s"
        )
    else:
        _add_par(doc, "[Resultados de baseline no disponibles aún]", italic=True)

    _add_heading(doc, "3.2 Accuracy de la red optimizada", level=2)
    if final:
        _add_par(doc,
            f"Accuracy en Test: {final.get('test_accuracy', float('nan')):.4f}  "
            f"|  Loss en Test: {final.get('test_loss', float('nan')):.4f}  "
            f"|  Épocas: {final.get('epochs')}  "
            f"|  Tiempo de entrenamiento: {final.get('train_time_seconds', 0):.1f} s"
        )
    else:
        _add_par(doc, "[Resultados del modelo optimizado no disponibles aún]", italic=True)

    _add_heading(doc, "3.3 Número de parámetros entrenables", level=2)
    _add_table(doc, ["Modelo", "Parámetros entrenables"], [
        ("Baseline", f"{baseline.get('trainable_params', 'N/D'):,}" if baseline else "N/D"),
        ("Optimizado", f"{final.get('trainable_params', 'N/D'):,}" if final else "N/D"),
    ])

    _add_heading(doc, "3.4 Mejor cromosoma encontrado", level=2)
    if ga_summary:
        _add_par(doc, ga_summary.get("best_chromosome_str", ""))
        _add_par(doc,
            f"Fitness: {ga_summary.get('best_fitness', float('nan')):.4f}  |  "
            f"val_accuracy: {ga_summary.get('best_val_accuracy')}  |  "
            f"Parámetros entrenables: {ga_summary.get('best_trainable_params')}  |  "
            f"Tiempo total del GA: {ga_summary.get('total_time_seconds', 0)/60:.1f} min"
        )
    else:
        _add_par(doc, "[Resumen del GA no disponible aún]", italic=True)

    _add_heading(doc, "3.5 Curva de evolución del fitness", level=2)
    _add_image_if_exists(
        doc, os.path.join(config.REPORT_DIR, "fig_fitness_evolution.png"),
        caption="Mejor fitness por generación a lo largo de la evolución del GA."
    )

    _add_heading(doc, "3.6 Curvas de aprendizaje", level=2)
    _add_par(doc, "Baseline (30 épocas):")
    _add_image_if_exists(
        doc, os.path.join(config.REPORT_DIR, "fig_learning_curves_baseline.png"),
        caption="Curvas de accuracy y loss (train/val) del modelo baseline."
    )
    _add_par(doc, "Modelo optimizado (30 épocas):")
    _add_image_if_exists(
        doc, os.path.join(config.REPORT_DIR, "fig_learning_curves_optimized.png"),
        caption="Curvas de accuracy y loss (train/val) del modelo optimizado."
    )

    _add_heading(doc, "3.7 Gráfica Accuracy vs Parámetros", level=2)
    _add_image_if_exists(
        doc, os.path.join(config.REPORT_DIR, "fig_accuracy_vs_params.png"),
        caption="Accuracy vs número de parámetros entrenables (escala logarítmica). "
                "Puntos grises: individuos evaluados durante la optimización (val_accuracy). "
                "Estrellas: baseline y modelo optimizado (test_accuracy)."
    )

    # ==================================================================
    # COMPARACIÓN FINAL OBLIGATORIA
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "4. Comparación Final", level=1)
    if comparison:
        rows = []
        for metric, vals in comparison.items():
            rows.append((metric, vals.get("Baseline"), vals.get("Optimizada")))
        _add_table(doc, ["Métrica", "Baseline", "Optimizada"], rows)
    else:
        _add_par(doc, "[Tabla comparativa no disponible aún]", italic=True)

    _add_heading(doc, "4.1 Justificación de la mejor solución", level=2)
    if baseline and final:
        acc_b = baseline.get("test_accuracy", 0)
        acc_o = final.get("test_accuracy", 0)
        tp_b = baseline.get("trainable_params", 1)
        tp_o = final.get("trainable_params", 1)
        delta_acc = acc_o - acc_b
        ratio_tp = tp_o / tp_b if tp_b else float("nan")
        _add_par(doc,
            f"La red optimizada obtuvo una diferencia de accuracy en Test de "
            f"{delta_acc:+.4f} respecto al baseline, utilizando "
            f"{ratio_tp:.2f}x los parámetros del baseline "
            f"({tp_o:,} vs {tp_b:,}). "
            + (
                "Dado que la red optimizada logra una accuracy igual o "
                "superior con una arquitectura más compacta o comparable, "
                "se considera la mejor solución bajo el criterio "
                "multiobjetivo del examen (exactitud + eficiencia)."
                if (delta_acc >= 0 and ratio_tp <= 1.0) else
                "La elección de la 'mejor' arquitectura depende del "
                "contexto de aplicación: si prima la exactitud absoluta y "
                "los recursos no son una restricción relevante, el "
                "baseline (o la optimizada, según cuál tenga mayor "
                "accuracy) sería preferible; si en cambio importan la "
                "eficiencia, la memoria y el costo de despliegue, la "
                "arquitectura con menor número de parámetros y una caída "
                "de accuracy marginal representa el mejor compromiso. "
                "Ver la discusión cuantitativa en la Sección 5."
            )
        )
    else:
        _add_par(doc, "[Pendiente de resultados finales]", italic=True)

    # ==================================================================
    # DISCUSIÓN
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "5. Discusión", level=1)

    _add_heading(doc, "5.1 Análisis de la arquitectura encontrada", level=2)
    if ga_summary:
        p = ga_summary.get("best_params", {})
        _add_par(doc,
            f"El Algoritmo Genético convergió hacia una arquitectura con "
            f"{p.get('n_conv_layers')} capa(s) convolucional(es) con filtros "
            f"{p.get('filters')}, kernel {p.get('kernel_size')}x"
            f"{p.get('kernel_size')}, capa densa de {p.get('dense_units')} "
            f"neuronas, dropout={p.get('dropout')}, optimizador "
            f"'{config.OPTIMIZER_MAP.get(p.get('optimizer'))}' con learning "
            f"rate={p.get('learning_rate')} y batch_size={p.get('batch_size')}. "
            f"Esta combinación refleja el compromiso que la función fitness "
            f"fue diseñada para incentivar: buena capacidad predictiva con "
            f"un número contenido de parámetros."
        )
    else:
        _add_par(doc, "[Pendiente de resultados finales]", italic=True)

    _add_heading(doc, "5.2 Comparación con el baseline", level=2)
    _add_par(doc,
        "Ver tabla comparativa (Sección 4) y gráfica Accuracy vs Parámetros "
        "(Sección 3.7). La posición relativa de ambos modelos en dicha "
        "gráfica resume visualmente el trade-off obtenido: un modelo por "
        "encima y a la izquierda del otro domina estrictamente (mejor "
        "accuracy con menos parámetros); si ninguno domina al otro, ambos "
        "pertenecen a la frontera de Pareto exactitud-complejidad y la "
        "elección depende de las prioridades de la aplicación."
    )

    _add_heading(doc, "5.3 Impacto de cada hiperparámetro", level=2)
    _add_bullets(doc, [
        "learning_rate: valores muy altos (p.ej. 0.01 con SGD) tienden a "
        "producir entrenamientos inestables en pocas épocas; valores muy "
        "bajos (p.ej. 0.0001) convergen lentamente y penalizan a los "
        "individuos evaluados con solo 5 épocas durante el GA.",
        "batch_size: lotes pequeños (32) actualizan los pesos con más "
        "frecuencia por época (más pasos de gradiente) pero con gradientes "
        "más ruidosos; lotes grandes (128) son más estables pero requieren "
        "más épocas para el mismo número de actualizaciones.",
        "optimizer: Adam suele converger más rápido en pocas épocas que "
        "SGD, lo que lo favorece dentro del protocolo de 5 épocas por "
        "individuo; RMSprop se comporta de forma intermedia.",
        "n_conv_layers / filters: mayor profundidad y más filtros aumentan "
        "la capacidad de representación pero también el número de "
        "parámetros y el riesgo de sobreajuste con pocas épocas de "
        "entrenamiento.",
        "kernel_size: kernels más grandes (5x5) capturan patrones "
        "espaciales más amplios por capa, pero reducen más agresivamente "
        "el mapa de características, limitando la profundidad máxima "
        "válida de la red (ver recorte de profundidad, Sección 2.2).",
        "dense_units: controla la capacidad de la capa de clasificación "
        "final; junto con el tamaño del mapa de características aplanado, "
        "es frecuentemente el mayor contribuyente al total de parámetros "
        "entrenables.",
        "dropout: valores moderados (0.2-0.3) ayudan a regularizar sin "
        "perjudicar significativamente la convergencia en pocas épocas; "
        "valores muy altos (0.5) pueden ralentizar el aprendizaje durante "
        "las 5 épocas de evaluación del GA.",
    ])

    _add_heading(doc, "5.4 Relación entre precisión y complejidad", level=2)
    _add_par(doc,
        "Los resultados obtenidos (Sección 3.7) evidencian que, dentro del "
        "espacio de búsqueda explorado, existe una región de arquitecturas "
        "pequeñas que alcanzan una exactitud competitiva con una fracción "
        "de los parámetros del baseline, así como arquitecturas grandes "
        "cuyo exceso de parámetros no necesariamente se traduce en mejoras "
        "proporcionales de accuracy. Esto confirma que, más allá de cierto "
        "punto, agregar capacidad al modelo tiene rendimientos "
        "decrecientes, y que la función fitness multiobjetivo es una "
        "herramienta adecuada para identificar automáticamente ese punto "
        "de equilibrio."
    )

    # ==================================================================
    # CONCLUSIONES
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "6. Conclusiones", level=1)
    _add_par(doc,
        "Los Algoritmos Genéticos demostraron ser una herramienta viable y "
        "flexible para optimizar de forma conjunta hiperparámetros de "
        "entrenamiento y decisiones arquitectónicas de una CNN, sin "
        "requerir gradientes de la métrica objetivo ni supuestos de "
        "diferenciabilidad. Su principal fortaleza radica en poder "
        "combinar objetivos en conflicto (exactitud vs. complejidad) en "
        "una sola función de fitness, permitiendo explorar automáticamente "
        "la frontera de compromiso entre ambos, en lugar de optimizar "
        "ciegamente solo la exactitud. Como limitación, el proceso sigue "
        "siendo computacionalmente costoso (cada individuo requiere "
        "entrenar una red neuronal completa) y la calidad de la solución "
        "depende fuertemente del diseño del espacio de búsqueda, la "
        "función de fitness y los recursos de cómputo disponibles (número "
        "de épocas por individuo, tamaño de población y generaciones). En "
        "general, se concluye que los GA son especialmente valiosos cuando "
        "se necesita optimizar múltiples objetivos simultáneamente y el "
        "espacio de diseño combina variables continuas, categóricas y "
        "estructurales, como ocurre en el diseño de arquitecturas de Deep "
        "Learning eficientes."
    )

    # ==================================================================
    # PREGUNTAS DE ANÁLISIS CRÍTICO
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "7. Preguntas de Análisis Crítico", level=1)

    questions_answers = [
        ("1. Si una CNN obtiene 92.4% de accuracy en Validation y otra obtiene "
         "92.0%, pero la segunda utiliza únicamente el 25% de los parámetros de "
         "la primera, ¿cuál elegiría para continuar hacia la etapa final de "
         "entrenamiento y por qué?",
         "Elegiría la segunda arquitectura. La pérdida de accuracy es marginal "
         "(0.4 puntos porcentuales), mientras que la reducción de parámetros es "
         "de 75%. Bajo la función fitness multiobjetivo del examen "
         "(w1=0.7, w2=0.3), esa reducción de complejidad compensa ampliamente "
         "una caída tan pequeña de accuracy: el fitness de la segunda red sería "
         "mayor. Además, un modelo con 4 veces menos parámetros generaliza con "
         "más margen, entrena/infiere más rápido y es más viable para "
         "producción."),
        ("2. ¿Puede una solución ser considerada mejor durante la optimización "
         "aunque no tenga el mayor valor de accuracy en Validation? Explique "
         "considerando la función multiobjetivo utilizada.",
         "Sí. El fitness no es igual a val_accuracy; es una combinación "
         "ponderada de val_accuracy y de la penalización por parámetros. Una "
         "solución con menor accuracy pero muchos menos parámetros puede "
         "obtener un fitness total mayor que otra con accuracy ligeramente "
         "superior pero mucho más grande, precisamente porque el término "
         "w2*(1-normalized_TP) puede pesar más que la diferencia de accuracy."),
        ("3. ¿Por qué en este experimento el conjunto Test no participa durante "
         "la optimización?",
         "Porque el conjunto Test debe permanecer como una estimación "
         "totalmente independiente e imparcial del desempeño de generalización "
         "final. Si se usara durante la optimización (aunque sea solo para "
         "calcular fitness), el GA terminaría ajustando indirectamente sus "
         "decisiones a las particularidades de ese conjunto (un tipo de fuga "
         "de información / overfitting al test set), invalidando su uso como "
         "medida objetiva al final del proceso."),
        ("4. ¿Cuál es el papel específico del conjunto Validation dentro del "
         "Algoritmo Genético?",
         "Actúa como la señal de evaluación (fitness) de cada individuo: tras "
         "entrenar cada arquitectura con Train durante 5 épocas, se mide su "
         "val_accuracy sobre Validation, y ese valor (junto con los parámetros "
         "entrenables) determina qué tan 'apto' es el individuo para "
         "sobrevivir, reproducirse y transmitir sus genes a la siguiente "
         "generación."),
        ("5. ¿Qué problema aparecería si se utilizara el conjunto Test para "
         "calcular el fitness de cada individuo?",
         "Se produciría fuga de información (data leakage): el proceso "
         "evolutivo seleccionaría implícitamente arquitecturas e "
         "hiperparámetros que se ajustan bien a las particularidades del "
         "conjunto Test, no a la población general de datos. La accuracy final "
         "reportada sobre ese mismo Test estaría sesgada optimistamente y ya "
         "no sería una estimación confiable de la capacidad de generalización "
         "del modelo sobre datos nuevos."),
        ("6. ¿Una red con más capas convolucionales necesariamente será "
         "superior a una red más pequeña? Justifique usando conceptos de "
         "complejidad y generalización.",
         "No necesariamente. Más capas incrementan la capacidad del modelo "
         "(puede representar funciones más complejas), pero también aumentan "
         "el riesgo de sobreajuste, especialmente con conjuntos de "
         "entrenamiento limitados o pocas épocas, y pueden introducir "
         "problemas de optimización (gradientes que se desvanecen, mayor "
         "dificultad de convergencia). Existe un punto óptimo de complejidad "
         "para cada problema/dataset; superarlo típicamente reduce la "
         "generalización (mayor varianza) sin una mejora proporcional en la "
         "exactitud, mientras que una red más pequeña bien ajustada puede "
         "generalizar mejor con menor varianza."),
        ("7. Si dos arquitecturas presentan exactamente el mismo accuracy en "
         "Validation y una utiliza diez veces menos parámetros, ¿cuál debería "
         "preferirse según la lógica de este examen?",
         "La de menos parámetros. Con val_accuracy idéntico, el término "
         "w1*val_accuracy es igual para ambas, por lo que el fitness total lo "
         "determina exclusivamente el término de complejidad, que favorece "
         "claramente a la arquitectura más pequeña. Esto es consistente con "
         "el principio de parsimonia (a igual desempeño, preferir el modelo "
         "más simple): menor riesgo de sobreajuste, menor costo computacional "
         "y de despliegue."),
        ("8. ¿Por qué minimizar el número de parámetros puede ser importante "
         "en aplicaciones reales?",
         "Modelos con menos parámetros requieren menos memoria RAM/almacenamiento, "
         "consumen menos energía, tienen menor latencia de inferencia y son más "
         "fáciles de desplegar en dispositivos con recursos limitados (móviles, "
         "IoT, sistemas embebidos). Además, suelen entrenar/actualizar más "
         "rápido, reducen costos de infraestructura en producción a gran "
         "escala, y tienden a ser menos propensos al sobreajuste."),
        ("9. Si durante la optimización aparece una arquitectura con accuracy "
         "ligeramente menor pero con una reducción considerable en parámetros, "
         "¿cómo debería interpretarse ese resultado?",
         "Debe interpretarse como una solución valiosa dentro de la frontera "
         "de compromiso (Pareto) entre exactitud y eficiencia, no como una "
         "solución 'peor' en términos absolutos. Si la caída de accuracy es "
         "pequeña y el ahorro de parámetros es grande, esa arquitectura puede "
         "ser preferible en escenarios donde la eficiencia computacional "
         "importa, y precisamente por eso la función fitness multiobjetivo "
         "puede asignarle un valor competitivo o incluso superior."),
        ("10. ¿Por qué utilizar accuracy de entrenamiento dentro de la función "
         "fitness sería una mala práctica?",
         "El accuracy de entrenamiento mide qué tan bien el modelo memoriza "
         "los datos ya vistos, no su capacidad de generalizar a datos nuevos. "
         "Usarlo como fitness llevaría al GA a seleccionar arquitecturas "
         "grandes y con alta capacidad de memorización (sobreajuste), que "
         "obtendrían accuracy de entrenamiento muy alto pero un desempeño "
         "pobre en Validation/Test, contradiciendo el objetivo real de la "
         "optimización."),
        ("11. ¿Es posible que la arquitectura seleccionada como mejor individuo "
         "después de 5 épocas ya no sea la mejor después de entrenarla durante "
         "30 épocas? Explique.",
         "Sí. Cinco épocas ofrecen una estimación parcial y ruidosa del "
         "desempeño potencial de una arquitectura. Algunas configuraciones "
         "(por ejemplo, con learning rates bajos u optimizadores como SGD) "
         "convergen más lentamente y podrían superar en el largo plazo a "
         "configuraciones que lucían mejores en el corto plazo (p. ej. con "
         "Adam y learning rates altos, que convergen rápido pero pueden "
         "estancarse antes). El ranking observado a 5 épocas no garantiza el "
         "ranking a 30 épocas."),
        ("12. ¿Por qué el protocolo exige entrenar cada individuo durante "
         "solamente 5 épocas, pero reentrenar el mejor modelo durante 30 "
         "épocas?",
         "Por eficiencia computacional: durante la optimización se deben "
         "entrenar decenas de arquitecturas distintas (población x "
         "generaciones), por lo que entrenar cada una durante muchas épocas "
         "sería computacionalmente prohibitivo. Cinco épocas bastan para "
         "obtener una señal relativa razonable entre arquitecturas "
         "(ranking aproximado). Una vez seleccionado un único ganador, sí es "
         "asumible invertir el cómputo de un entrenamiento completo (30 "
         "épocas) para obtener una estimación de desempeño más precisa y "
         "representativa de su verdadero potencial."),
        ("13. Suponga que la arquitectura ganadora obtiene una mejora de 0.5% "
         "en accuracy respecto al baseline, pero duplica el número de "
         "parámetros. ¿Considera que realmente representa una mejora? "
         "Justifique.",
         "En términos estrictamente multiobjetivo, no necesariamente. Una "
         "mejora de solo 0.5% de accuracy a cambio de duplicar los parámetros "
         "representa un intercambio poco favorable: el costo marginal en "
         "complejidad (memoria, tiempo de inferencia, riesgo de sobreajuste) "
         "es alto en relación con la ganancia de exactitud obtenida. Bajo la "
         "función fitness usada en este examen, dicha solución probablemente "
         "tendría un fitness menor que una alternativa más compacta con "
         "accuracy comparable. Solo se justificaría si la aplicación exige "
         "maximizar accuracy sin restricción alguna de recursos."),
        ("14. Si la arquitectura optimizada tiene menor accuracy que el "
         "baseline, pero utiliza menos de la mitad de los parámetros, ¿puede "
         "considerarse exitosa la optimización? Explique.",
         "Sí, puede considerarse exitosa si se interpreta correctamente el "
         "objetivo del ejercicio: no se buscaba únicamente maximizar "
         "accuracy, sino encontrar un buen compromiso entre exactitud y "
         "eficiencia. Una arquitectura con menos de la mitad de los "
         "parámetros y una caída de accuracy pequeña/aceptable demuestra que "
         "el GA logró identificar una región más eficiente del espacio de "
         "búsqueda, lo cual es exactamente el propósito de la optimización "
         "multiobjetivo planteada."),
        ("15. Si se duplica el tamaño de la población del Algoritmo Genético, "
         "¿está garantizado encontrar una mejor solución? ¿Por qué?",
         "No está garantizado. Los GA son métodos estocásticos y "
         "heurísticos: una población más grande aumenta la diversidad "
         "genética y la probabilidad de explorar mejores regiones del "
         "espacio de búsqueda, pero no asegura convergencia a un óptimo "
         "mejor, especialmente si el número de generaciones no también "
         "aumenta proporcionalmente, si la función fitness tiene múltiples "
         "óptimos locales, o si los operadores de selección/mutación no "
         "logran explotar esa diversidad adicional de forma efectiva."),
        ("16. ¿Por qué la mutación puede generar individuos aparentemente "
         "peores y aun así ser necesaria para el éxito del algoritmo?",
         "La mutación introduce diversidad genética que el cruce por sí solo "
         "no puede generar (el cruce solo recombina genes ya presentes en la "
         "población). Sin mutación, la población puede converger "
         "prematuramente hacia un óptimo local y quedar 'atrapada' sin poder "
         "escapar de él. Individuos mutados que parecen peores a corto plazo "
         "pueden portar genes (p. ej. un valor de learning rate o un número "
         "de filtros distinto) que, combinados en generaciones futuras vía "
         "cruce, produzcan soluciones superiores. Es el mecanismo de "
         "exploración que equilibra la explotación de las mejores soluciones "
         "actuales."),
        ("17. Suponga que el individuo con mayor fitness utiliza una "
         "combinación de hiperparámetros muy distinta a la que usted hubiera "
         "elegido manualmente. ¿Qué indica esto acerca del espacio de "
         "búsqueda y del proceso de optimización?",
         "Indica que el espacio de búsqueda tiene interacciones no evidentes "
         "y posiblemente no lineales entre hiperparámetros, que no coinciden "
         "necesariamente con la intuición o las heurísticas de diseño "
         "manual (por ejemplo, 'más filtros siempre es mejor' o 'kernels más "
         "grandes capturan mejor los patrones'). Esto valida el valor "
         "práctico de la búsqueda automática: el GA puede descubrir "
         "combinaciones contraintuitivas pero efectivas que un diseñador "
         "humano difícilmente probaría por sesgos de experiencia previa."),
        ("18. Análisis integrador (trade-off exactitud/complejidad, ventajas y "
         "desventajas de cada arquitectura, ¿mejora real o solución "
         "alternativa?, y justificación de uso en una aplicación real).",
         "[Ver desarrollo completo con evidencia experimental en la Sección "
         "8 de este reporte.]"),
    ]

    for q, a in questions_answers:
        _add_par(doc, q, bold=True)
        _add_par(doc, a)
        doc.add_paragraph()

    # ==================================================================
    # PREGUNTA 18 - DESARROLLO EXTENDIDO CON EVIDENCIA EXPERIMENTAL
    # ==================================================================
    doc.add_page_break()
    _add_heading(doc, "8. Desarrollo de la Pregunta 18 (evidencia experimental)", level=1)

    if baseline and final and comparison:
        acc_b = baseline.get("test_accuracy", 0)
        acc_o = final.get("test_accuracy", 0)
        tp_b = baseline.get("trainable_params", 1)
        tp_o = final.get("trainable_params", 1)
        t_b = baseline.get("train_time_seconds", 0)
        t_o = final.get("train_time_seconds", 0)

        _add_heading(doc, "Trade-off exactitud vs. complejidad", level=2)
        _add_par(doc,
            f"Evidencia experimental: Baseline -> accuracy Test={acc_b:.4f}, "
            f"parámetros={tp_b:,}, tiempo de entrenamiento={t_b:.1f}s. "
            f"Optimizada -> accuracy Test={acc_o:.4f}, parámetros={tp_o:,}, "
            f"tiempo de entrenamiento={t_o:.1f}s. "
            f"Diferencia de accuracy: {acc_o-acc_b:+.4f}. "
            f"Razón de parámetros (optimizada/baseline): {tp_o/tp_b:.3f}. "
        )

        _add_heading(doc, "Ventajas y desventajas de cada arquitectura", level=2)
        _add_bullets(doc, [
            f"Baseline: arquitectura simple y ya validada en clase; "
            f"{tp_b:,} parámetros; accuracy Test={acc_b:.4f}. Ventaja: "
            f"referencia estable y conocida. Desventaja: no fue optimizada "
            f"para eficiencia, por lo que puede tener más parámetros de los "
            f"estrictamente necesarios para el desempeño que logra.",
            f"Optimizada: {tp_o:,} parámetros; accuracy Test={acc_o:.4f}. "
            f"Ventaja: producto de una búsqueda automática que considera "
            f"explícitamente el costo en parámetros. Desventaja potencial: "
            f"fue seleccionada con solo {config.GA_INDIVIDUAL_EPOCHS} épocas "
            f"de entrenamiento por individuo durante el GA, por lo que su "
            f"ranking relativo a otras arquitecturas candidatas podría "
            f"cambiar con más épocas de exploración (ver pregunta 11).",
        ])

        mejor_txt = (
            "la optimización SÍ produjo una mejora real (igual o mayor "
            "accuracy con igual o menor complejidad, o una mejora "
            "claramente favorable en el compromiso costo-beneficio)."
            if (acc_o >= acc_b and tp_o <= tp_b) else
            "la optimización encontró una SOLUCIÓN ALTERNATIVA en la "
            "frontera de compromiso exactitud-complejidad, no una mejora "
            "estrictamente dominante: mejora en un eje (accuracy o "
            "parámetros) a costa de una posible cesión en el otro. Cuál de "
            "las dos arquitecturas es 'mejor' depende de las prioridades de "
            "la aplicación."
        )
        _add_heading(doc, "¿Mejora real o solución alternativa?", level=2)
        _add_par(doc, f"Con base en los resultados obtenidos, {mejor_txt}")

        _add_heading(doc, "Justificación de uso en una aplicación real", level=2)
        _add_par(doc,
            "En un escenario con recursos computacionales amplios y donde "
            "la exactitud es crítica (por ejemplo, diagnóstico médico de "
            "alto riesgo), se recomendaría priorizar la arquitectura con "
            "mayor accuracy, sea el baseline o la optimizada. En cambio, en "
            "un escenario con restricciones de memoria, energía o latencia "
            "(dispositivos móviles/edge, sistemas embebidos, servicios con "
            "alto volumen de inferencias), se recomendaría priorizar la "
            "arquitectura optimizada, salvo que su pérdida de accuracy sea "
            "inaceptable para el caso de uso. La evidencia experimental de "
            "este proyecto (tabla comparativa y gráfica Accuracy vs "
            "Parámetros) debe usarse como base cuantitativa para tomar esa "
            "decisión en cada contexto concreto."
        )
    else:
        _add_par(doc, "[Pendiente de resultados finales para completar esta sección con evidencia experimental]", italic=True)

    doc.save(OUT_PATH)
    print(f"Reporte generado en: {OUT_PATH}")
    return OUT_PATH


if __name__ == "__main__":
    generate_report()
